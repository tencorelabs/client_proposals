#!/usr/bin/env python3
"""
TenCore Digital Labs • High-Fidelity Word (.docx) Proposal Generator
Parses Markdown proposals and generates beautifully styled, executive-grade Word documents.
"""

import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# --- BRAND PALETTE ---
COLOR_PRIMARY_NAVY = RGBColor(15, 41, 66)      # #0F2942 (Deep Navy)
COLOR_ACCENT_BLUE   = RGBColor(2, 132, 199)    # #0284C7 (TenCore Blue)
COLOR_TEXT_DARK     = RGBColor(30, 41, 59)     # #1E293B (Slate 800)
COLOR_MUTED_GRAY    = RGBColor(100, 116, 139)  # #64748B (Slate 500)
HEX_NAVY_BG         = "0F2942"
HEX_ZEBRA_BG        = "F8FAFC"
HEX_WHITE_BG        = "FFFFFF"
HEX_CALLOUT_BG      = "F0F9FF"
HEX_CODE_BG         = "F1F5F9"
HEX_BORDER_LIGHT    = "CBD5E1"
HEX_ACCENT_BLUE     = "0284C7"

def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets padding inside table cells (in dxa: 20 dxa = 1 pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    """Applies subtle borders to a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def set_callout_border(cell, border_color="0284C7"):
    """Adds a left thick accent border to a callout cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
            <w:top w:val="none"/>
            <w:right w:val="none"/>
            <w:bottom w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

def add_formatted_text(paragraph, text, default_color=COLOR_TEXT_DARK, default_size=Pt(10), default_bold=False, default_italic=False, default_font="Calibri"):
    """Parses basic inline Markdown (**bold**, *italic*, `code`) into docx runs."""
    # Pattern to match **bold**, *italic*, `code`
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run()
        run.font.name = default_font
        run.font.size = default_size
        run.font.color.rgb = default_color
        run.bold = default_bold
        run.italic = default_italic
        
        if token.startswith('**') and token.endswith('**'):
            run.text = token[2:-2]
            run.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run.text = token[1:-1]
            run.italic = True
        elif token.startswith('`') and token.endswith('`'):
            run.text = token[1:-1]
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = COLOR_PRIMARY_NAVY
        else:
            run.text = token

def build_docx_proposal(md_path, docx_out_path, logo_path=None):
    doc = Document()
    
    # Page Margins (0.85 in for clean spacious margins)
    for section in doc.sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)
        
        # Header setup
        header = section.header
        header_p = header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = header_p.add_run("TENCORE DIGITAL LABS • CLIENT COMMERCIAL PROPOSAL")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8)
        hrun.font.color.rgb = COLOR_MUTED_GRAY
        
        # Footer setup
        footer = section.footer
        footer_p = footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = footer_p.add_run("CONFIDENTIAL • TenCore Digital Labs Private Limited © 2026")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8)
        frun.font.color.rgb = COLOR_MUTED_GRAY

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    # --- TITLE & METADATA SECTION ---
    # Add Logo if available
    if logo_path and os.path.exists(logo_path):
        logo_p = doc.add_paragraph()
        logo_p.paragraph_format.space_before = Pt(0)
        logo_p.paragraph_format.space_after = Pt(12)
        logo_p.add_run().add_picture(logo_path, width=Inches(2.2))

    # Read Title & Subtitle from first few lines
    title_text = "Digital Valet Parking System"
    subtitle_text = "MVP Scope, Technology Architecture, Cost & Commercial Terms"
    meta_info = []

    idx = 0
    in_header_meta = True
    while idx < len(lines):
        line = lines[idx].strip()
        if line.startswith('# '):
            title_text = line[2:].strip()
            idx += 1
            continue
        if line.startswith('## '):
            if in_header_meta:
                subtitle_text = line[3:].strip()
                idx += 1
                continue
        if line.startswith('---') and in_header_meta:
            in_header_meta = False
            idx += 1
            break
        if in_header_meta and line.startswith('**'):
            # metadata line
            meta_info.append(line)
        idx += 1

    # Document Header Box
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run(title_text)
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY_NAVY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run(subtitle_text)
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = COLOR_ACCENT_BLUE

    # Metadata Card Table
    if meta_info:
        meta_table = doc.add_table(rows=1, cols=1)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_table.autofit = False
        cell = meta_table.rows[0].cells[0]
        cell.width = Inches(6.8)
        set_cell_background(cell, HEX_CODE_BG)
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
        set_callout_border(cell, border_color=HEX_NAVY_BG)
        
        for m_line in meta_info:
            mp = cell.add_paragraph()
            mp.paragraph_format.space_before = Pt(1)
            mp.paragraph_format.space_after = Pt(1)
            mp.paragraph_format.line_spacing = 1.15
            add_formatted_text(mp, m_line, default_size=Pt(9.5), default_color=COLOR_TEXT_DARK)

    # Space after header block
    div_p = doc.add_paragraph()
    div_p.paragraph_format.space_before = Pt(10)
    div_p.paragraph_format.space_after = Pt(6)

    # --- PARSING BODY CONTENT ---
    in_code_block = False
    code_lines = []
    
    in_table = False
    table_lines = []

    def flush_code_block(lines):
        if not lines:
            return
        code_text = "\n".join(lines)
        ctable = doc.add_table(rows=1, cols=1)
        ctable.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = ctable.rows[0].cells[0]
        cell.width = Inches(6.8)
        set_cell_background(cell, HEX_CODE_BG)
        set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
        set_callout_border(cell, border_color=HEX_BORDER_LIGHT)
        cp = cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(0)
        cp.paragraph_format.space_after = Pt(0)
        cp.paragraph_format.line_spacing = 1.1
        crun = cp.add_run(code_text)
        crun.font.name = "Consolas"
        crun.font.size = Pt(8.5)
        crun.font.color.rgb = COLOR_TEXT_DARK
        
        # Spacer
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(4)

    def flush_table(raw_lines):
        if not raw_lines:
            return
        parsed_rows = []
        for rl in raw_lines:
            # Split by |
            parts = [c.strip() for c in rl.strip().split('|')]
            if len(parts) >= 3 and parts[0] == '' and parts[-1] == '':
                parts = parts[1:-1]
            elif parts and parts[0] == '':
                parts = parts[1:]
            if parts and parts[-1] == '':
                parts = parts[:-1]
            # Check if divider line like |---|---|
            if all(set(c).issubset({'-', ':', ' '}) for c in parts):
                continue
            parsed_rows.append(parts)
        
        if not parsed_rows:
            return
        
        num_cols = max(len(r) for r in parsed_rows)
        # Pad shorter rows
        for r in parsed_rows:
            while len(r) < num_cols:
                r.append("")
                
        t = doc.add_table(rows=len(parsed_rows), cols=num_cols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(t, color=HEX_BORDER_LIGHT, sz="4")
        
        # Distribute column widths
        total_width = 6.8
        col_width = total_width / num_cols
        
        for r_idx, row_data in enumerate(parsed_rows):
            is_header = (r_idx == 0)
            row = t.rows[r_idx]
            
            # Prevent row split across pages
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            
            if is_header:
                trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            
            for c_idx, val in enumerate(row_data):
                cell = row.cells[c_idx]
                cell.width = Inches(col_width)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
                
                # Check alignment (right align currency / percentages / numbers)
                align = WD_ALIGN_PARAGRAPH.LEFT
                clean_val = val.replace('₹', '').replace(',', '').replace('%', '').strip()
                if any(val.startswith(x) for x in ['₹', '$', '30%', '25%', '15%', '100%']) or clean_val.replace('.', '', 1).isdigit():
                    if not is_header:
                        align = WD_ALIGN_PARAGRAPH.RIGHT
                
                p = cell.paragraphs[0]
                p.alignment = align
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                
                if is_header:
                    set_cell_background(cell, HEX_NAVY_BG)
                    add_formatted_text(p, val, default_color=RGBColor(255, 255, 255), default_size=Pt(9), default_bold=True)
                else:
                    bg_color = HEX_ZEBRA_BG if (r_idx % 2 == 1) else HEX_WHITE_BG
                    set_cell_background(cell, bg_color)
                    add_formatted_text(p, val, default_size=Pt(9), default_color=COLOR_TEXT_DARK)
        
        # Spacer after table
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(0)
        sp.paragraph_format.space_after = Pt(6)

    # Process remaining lines
    while idx < len(lines):
        line = lines[idx]
        s_line = line.strip()

        # Handle Code blocks
        if s_line.startswith('```'):
            if in_code_block:
                flush_code_block(code_lines)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
            idx += 1
            continue

        if in_code_block:
            code_lines.append(line.rstrip('\r\n'))
            idx += 1
            continue

        # Handle Tables
        if '|' in s_line and s_line.startswith('|') and s_line.endswith('|'):
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(s_line)
            idx += 1
            continue
        else:
            if in_table:
                flush_table(table_lines)
                table_lines = []
                in_table = False

        # Divider lines
        if s_line.startswith('---'):
            idx += 1
            continue

        # Headings
        if s_line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            add_formatted_text(p, s_line[2:], default_color=COLOR_PRIMARY_NAVY, default_size=Pt(16), default_bold=True)
            idx += 1
            continue

        if s_line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            add_formatted_text(p, s_line[3:], default_color=COLOR_PRIMARY_NAVY, default_size=Pt(13), default_bold=True)
            idx += 1
            continue

        if s_line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            add_formatted_text(p, s_line[4:], default_color=COLOR_ACCENT_BLUE, default_size=Pt(11), default_bold=True)
            idx += 1
            continue

        # Blockquote / Callout
        if s_line.startswith('> '):
            quote_text = s_line[2:].strip()
            q_table = doc.add_table(rows=1, cols=1)
            q_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = q_table.rows[0].cells[0]
            cell.width = Inches(6.8)
            set_cell_background(cell, HEX_CALLOUT_BG)
            set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
            set_callout_border(cell, border_color=HEX_ACCENT_BLUE)
            qp = cell.paragraphs[0]
            qp.paragraph_format.space_before = Pt(2)
            qp.paragraph_format.space_after = Pt(2)
            qp.paragraph_format.line_spacing = 1.15
            add_formatted_text(qp, quote_text, default_size=Pt(9.5), default_color=COLOR_PRIMARY_NAVY, default_italic=True)
            
            sp = doc.add_paragraph()
            sp.paragraph_format.space_before = Pt(0)
            sp.paragraph_format.space_after = Pt(4)
            idx += 1
            continue

        # Unordered Lists
        if s_line.startswith('- ') or s_line.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_text(p, s_line[2:], default_size=Pt(9.5), default_color=COLOR_TEXT_DARK)
            idx += 1
            continue

        # Ordered Lists
        match_ol = re.match(r'^(\d+)\.\s+(.*)$', s_line)
        if match_ol:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_text(p, match_ol.group(2), default_size=Pt(9.5), default_color=COLOR_TEXT_DARK)
            idx += 1
            continue

        # Regular Body Paragraph
        if s_line:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_text(p, s_line, default_size=Pt(9.5), default_color=COLOR_TEXT_DARK)

        idx += 1

    # Flush any remaining tables
    if in_table:
        flush_table(table_lines)
    if in_code_block:
        flush_code_block(code_lines)

    doc.save(docx_out_path)
    print(f"✅ Generated Proposal Word Document: {docx_out_path}")

if __name__ == "__main__":
    md_file = "/Users/devendra.vaja/Development/TenCoreLabs/Client Proposals/ParkingSetu/References/digital_valet_parking_mvp_scope_and_cost.md"
    docx_file = "/Users/devendra.vaja/Development/TenCoreLabs/Client Proposals/ParkingSetu/ParkingSetu_MVP_Commercial_Proposal.docx"
    logo_file = "/Users/devendra.vaja/Development/TenCoreLabs/Client Proposals/ParkingSetu/Presentation/tencore_digital_labs_logo.png"
    
    build_docx_proposal(md_file, docx_file, logo_file)
